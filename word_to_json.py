"""
从 Word 文档中提取选择题并保存为 JSON 格式

将 test.docx 中的选择题表格解析出来，整理成 JSON 数组
"""

import argparse
import json
import re
from pathlib import Path

from docx import Document


def extract_questions_from_docx(docx_path: Path) -> list[dict]:
    """
    从 Word 文档中提取选择题
    
    Args:
        docx_path: Word 文档路径
        
    Returns:
        选择题列表，每题包含 title 和 options
    """
    document = Document(docx_path)
    questions = []
    
    for table in document.tables:
        try:
            # 期望的表格结构：
            # 第1行（合并单元格）：题目
            # 第2行：A选项 | B选项
            # 第3行：C选项 | D选项
            
            if len(table.rows) < 3:
                continue
            
            # 提取题目（第一行，合并的单元格）
            title_cell = table.cell(0, 0)
            title = title_cell.text.strip()
            
            # 如果题目为空，跳过这个表格
            if not title:
                continue
            
            # 提取四个选项
            options = {"a": "", "b": "", "c": "", "d": ""}
            
            # A 选项 - 第2行第1列
            a_text = table.cell(1, 0).text.strip()
            options["a"] = remove_option_prefix(a_text, "A")
            
            # B 选项 - 第2行第2列
            b_text = table.cell(1, 1).text.strip()
            options["b"] = remove_option_prefix(b_text, "B")
            
            # C 选项 - 第3行第1列
            c_text = table.cell(2, 0).text.strip()
            options["c"] = remove_option_prefix(c_text, "C")
            
            # D 选项 - 第3行第2列
            d_text = table.cell(2, 1).text.strip()
            options["d"] = remove_option_prefix(d_text, "D")
            
            questions.append({
                "title": title,
                "options": options
            })
            
        except Exception as e:
            print(f"警告: 解析表格时出错: {e}")
            continue
    
    return questions


def remove_option_prefix(text: str, letter: str) -> str:
    """
    移除选项前缀（如 "A. " 或 "A、"）
    
    Args:
        text: 选项文本
        letter: 选项字母 (A/B/C/D)
        
    Returns:
        去除前缀后的选项内容
    """
    # 匹配各种格式: A. / A、/ A: / A） / A) 等
    pattern = rf"^{letter}[\.\、\:\）\)\s]+\s*"
    return re.sub(pattern, "", text, flags=re.IGNORECASE).strip()


def main():
    parser = argparse.ArgumentParser(
        description="从 Word 文档中提取选择题并保存为 JSON"
    )
    parser.add_argument(
        "input",
        type=Path,
        help="输入的 Word 文档路径"
    )
    parser.add_argument(
        "-o", "--output",
        type=Path,
        default=None,
        help="输出 JSON 文件路径（默认为输入文件名.json）"
    )
    parser.add_argument(
        "--pretty",
        action="store_true",
        default=True,
        help="格式化 JSON 输出（默认启用）"
    )
    
    args = parser.parse_args()
    
    # 检查输入文件
    if not args.input.exists():
        print(f"错误: 文件不存在: {args.input}")
        return 1
    
    # 确定输出路径
    output_path = args.output
    if output_path is None:
        output_path = args.input.with_suffix(".json")
    
    try:
        # 提取选择题
        print(f"正在从 {args.input} 提取选择题...")
        questions = extract_questions_from_docx(args.input)
        
        if not questions:
            print("未提取到任何选择题")
            return 1
        
        print(f"成功提取 {len(questions)} 道选择题")
        
        # 保存为 JSON
        indent = 2 if args.pretty else None
        json_output = json.dumps(questions, ensure_ascii=False, indent=indent)
        output_path.write_text(json_output, encoding="utf-8")
        
        print(f"JSON 已保存到: {output_path}")
        
        # 打印预览
        print("\n提取结果预览:")
        for i, q in enumerate(questions[:3], 1):
            print(f"\n题目 {i}: {q['title'][:50]}...")
            print(f"  A: {q['options']['a'][:30]}...")
        
        if len(questions) > 3:
            print(f"\n... 共 {len(questions)} 道题目")
        
    except Exception as e:
        print(f"错误: {e}")
        return 1
    
    return 0


if __name__ == "__main__":
    exit(main())
